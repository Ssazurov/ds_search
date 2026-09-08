"""Экспорт глоссария и справочника ссылок (issue #25, ADR-005).

xlsx -> data/exports/*.json (для ds_site) + data/raw/{glossary,links}/
*.json+*.md (для ds_ingestion адаптера, issue #5).

Запуск: python -m scripts.export_glossary_links
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DS_ROOT = ROOT.parent  # ~/ds — где лежат исходные xlsx
GLOSSARY_XLSX = DS_ROOT / "data" / "downsyndrome_glossary.xlsx"
LINKS_XLSX = DS_ROOT / "data" / "sites_ru_down_syndrome.xlsx"
EXPORTS_DIR = ROOT / "data" / "exports"
RAW_DIR = ROOT / "data" / "raw"


def read_glossary(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Глоссарий СД"]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r[0] == "ID")
    out = []
    for r in rows[header_idx + 1:]:
        if not r[0] or not r[1]:
            continue
        out.append({
            "id": r[0], "term": r[1], "definition": r[2] or "",
            "en": r[3] or "", "category": r[4] or "",
        })
    return out


def read_links(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Сайты РФ о синдроме Дауна"]
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows[1:]:
        if not r[1] or not r[2]:
            continue
        out.append({
            "name": r[1], "url": r[2], "age": r[3] or "",
            "description": r[4] or "", "region": r[5] or "",
            "category": r[6] or "", "relevance": r[7] or "",
        })
    return out


def write_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def glossary_markdown(items: list[dict]) -> str:
    lines = ["# Глоссарий: синдром Дауна", ""]
    for it in items:
        lines.append(f"## {it['term']}" + (f" ({it['en']})" if it["en"] else ""))
        lines.append(it["definition"])
        if it["category"]:
            lines.append(f"*Категория: {it['category']}*")
        lines.append("")
    return "\n".join(lines)


def links_markdown(items: list[dict]) -> str:
    lines = ["# Справочник ресурсов о синдроме Дауна (РФ)", ""]
    for it in items:
        lines.append(f"## {it['name']}")
        lines.append(f"Сайт: {it['url']}")
        if it["description"]:
            lines.append(it["description"])
        meta = ", ".join(
            x for x in [it["region"], it["category"], it["age"] and f"возраст: {it['age']}"] if x
        )
        if meta:
            lines.append(f"*{meta}*")
        lines.append("")
    return "\n".join(lines)


def write_rag_doc(source_dir: Path, slug: str, title: str, content_md: str) -> None:
    source_dir.mkdir(parents=True, exist_ok=True)
    md_path = source_dir / f"{slug}.md"
    json_path = source_dir / f"{slug}.json"
    md_path.write_text(content_md, encoding="utf-8")
    meta = {
        "source_url": f"internal://ds_search/{slug}",
        "source_domain": "ds_search",
        "title": title,
        "direction": "methodology",
        "category": "inclusion",
        "doc_type": "glossary" if slug == "glossary" else "resource_directory",
        "license": "own_generated",
        "target_audience": "parents,specialists",
        "content_path": str(md_path),
    }
    write_json(json_path, meta)


def main() -> None:
    glossary = read_glossary(GLOSSARY_XLSX)
    links = read_links(LINKS_XLSX)

    write_json(EXPORTS_DIR / "glossary.json", glossary)
    write_json(EXPORTS_DIR / "links.json", links)

    write_rag_doc(
        RAW_DIR / "glossary", "glossary",
        "Глоссарий терминов: синдром Дауна", glossary_markdown(glossary),
    )
    write_rag_doc(
        RAW_DIR / "links", "links",
        "Справочник ресурсов о синдроме Дауна (РФ)", links_markdown(links),
    )
    print(f"glossary: {len(glossary)} terms, links: {len(links)} sites")


if __name__ == "__main__":
    main()
